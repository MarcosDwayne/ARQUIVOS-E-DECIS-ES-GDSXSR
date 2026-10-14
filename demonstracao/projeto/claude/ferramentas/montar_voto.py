#!/usr/bin/env python3
"""Monta a minuta de voto (APC) no padrão GDSXSR a partir de um JSON — v6 (06/10/2026).

Uso:
    python3 montar_voto.py minuta.json SAIDA.docx [--modelo=assets/modelo-base-apc.docx]

v6 (calibrações C-36 a C-62, MANUAL ATUALIZADO e estrutura pronta APC):
- texto normal com recuo de 1ª linha 4 cm; transcrições Arial 13 itálico com recuo 4 cm e 1ª linha 6 cm;
  ementas e lei transcritas (tipo "transcricao_ementa") com recuo 4 cm sem recuo de 1ª linha;
- subtítulo de capítulo realçado em amarelo; "veredito" ("Sem razão.") realçado; "mini_dispositivo"
  sublinhado e realçado (nunca no último capítulo);
- rótulos "Dispositivos relevantes citados:" / "Jurisprudência relevante citada:" em itálico;
- bloco de resumo em vermelho EE0000 entre a ementa e o relatório (resumo_titulo/resumo_secao/
  resumo_topico/resumo_texto);
- sentença anexada em verde 00B050.

Marcação inline: **negrito**, ==realce amarelo==, __sublinhado__.
"""
import json, re, sys, os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
MODELO_PADRAO = os.path.join(AQUI, "..", "assets", "modelo-base-apc.docx")
AUTOR_COMENTARIO = "ASSESSOR GDSXSR"
PREFIXO_COMENTARIO = "ASSESSOR GDSXSR — "
VERMELHO, VERDE = "EE0000", "00B050"

J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT

# tipo -> dict de formatação
B = dict(al=J, tam=14, neg=False, ita=False, sub=False, realce=False, cor=None, rec1=None, recE=None, vazio=True)
def E(**k):
    d = dict(B); d.update(k); return d
ESTILOS = {
    "secao":              E(al=C, neg=True),
    "ementa":             E(vazio=False),
    "ementa_titulo":      E(neg=True, vazio=False),
    "ementa_nota":        E(tam=13, vazio=False),
    "resumo_titulo":      E(neg=True, cor=VERMELHO, realce=True, vazio=False),
    "resumo_secao":       E(neg=True, cor=VERMELHO, vazio=False),
    "resumo_topico":      E(neg=True, cor=VERMELHO, vazio=False),
    "resumo_texto":       E(cor=VERMELHO, rec1=-0.63, recE=1.27, vazio=False),
    "paragrafo":          E(rec1=4.0),
    "veredito":           E(rec1=4.0, realce=True),
    "mini_dispositivo":   E(rec1=4.0, sub=True, realce=True),
    "transcricao":        E(tam=13, ita=True, rec1=2.0, recE=4.0, vazio=False),
    "transcricao_ementa": E(tam=13, ita=True, rec1=None, recE=4.0, vazio=False),
    "relator":            E(neg=True),
    "titulo_recurso":     E(al=C, neg=True),
    "subtitulo":          E(al=C, realce=True),
    "dispositivo":        E(rec1=4.0),
    "contraria_titulo":   E(al=C, neg=True, cor="800080"),
    "contraria":          E(tam=13, cor="800080", recE=4.0),
    "anexo_par":          E(tam=13.5, cor=VERDE, vazio=False),
}
TRANSCRICOES = ("transcricao", "transcricao_ementa")
TOKEN = re.compile(r"(\*\*.+?\*\*|==.+?==|(?<!_)__[^_].*?__(?!_))")
ROTULOS_ITALICO = ("Dispositivos relevantes citados:", "Jurisprudência relevante citada:")


def _run(p, texto, st, neg=None, realce=None, sub=None, ita=None):
    r = p.add_run(texto)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    r.font.size = Pt(st["tam"])
    r.bold = st["neg"] if neg is None else neg
    r.italic = st["ita"] if ita is None else ita
    if st["sub"] or sub:
        r.underline = True
    if st["realce"] or realce:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if st["cor"]:
        r.font.color.rgb = RGBColor.from_string(st["cor"])
    return r


def runs_marcados(p, texto, st):
    runs = []
    rot = next((x for x in ROTULOS_ITALICO if texto.startswith(x)), None)
    if rot:                                   # C-40: só o rótulo em itálico
        runs.append(_run(p, rot, st, ita=True))
        texto = texto[len(rot):]
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        kw = {}
        if parte.startswith("**") and parte.endswith("**"):
            parte, kw = parte[2:-2], {"neg": True}
        elif parte.startswith("==") and parte.endswith("=="):
            parte, kw = parte[2:-2], {"realce": True}
        elif parte.startswith("__") and parte.endswith("__"):
            parte, kw = parte[2:-2], {"sub": True}
        runs.append(_run(p, parte, st, **kw))
    return runs


def novo_par(doc, tipo, texto, quebra_antes=False):
    st = ESTILOS[tipo]
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = L if (tipo == "secao" and texto.strip().upper() == "EMENTA") else st["al"]
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    if st["rec1"] is not None:
        pf.first_line_indent = Cm(st["rec1"])
    if st["recE"] is not None:
        pf.left_indent = Cm(st["recE"])
    if quebra_antes:
        pf.page_break_before = True
    return p, runs_marcados(p, texto, st)


def vazio(doc):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(4.0)
    r = p.add_run("")
    r.font.size = Pt(14)
    return p


def preencher_autuacao(doc, aut):
    mapa = {"Processo": aut.get("processo", ""), "Apelante": aut.get("apelantes", ""),
            "Apelado": aut.get("apelados", ""), "Órgão": aut.get("orgao"),
            "Classe": aut.get("classe"), "Relator": aut.get("relator")}
    tab = doc.tables[1]._tbl
    for tr in tab.findall(qn("w:tr")):
        tcs = tr.findall(".//" + qn("w:tc"))
        rotulo = tcs[0].xpath("string(.)").strip()
        for chave, valor in mapa.items():
            if rotulo.startswith(chave) and valor:
                ts = tcs[2].findall(".//" + qn("w:t"))
                if not ts:
                    continue
                ts[0].text = valor
                for t in ts[1:]:
                    t.text = ""


def montar(spec, saida, modelo):
    doc = Document(modelo)
    preencher_autuacao(doc, spec.get("autuacao", {}))
    pendentes = []
    blocos = spec["blocos"]
    for i, b in enumerate(blocos):
        tipo = b["tipo"]
        prox = blocos[i + 1]["tipo"] if i + 1 < len(blocos) else None
        if tipo == "anexo":                     # C-45: decisão recorrida em verde, nova página
            p, _ = novo_par(doc, "anexo_par", b.get("titulo", "SENTENÇA"), quebra_antes=True)
            p.alignment = C
            for r in p.runs:
                r.bold = True
            for t in b.get("paragrafos", []):
                novo_par(doc, "anexo_par", t)
            continue
        if tipo == "resumo_titulo" and i > 0:
            vazio(doc); vazio(doc)
        quebra = tipo == "contraria_titulo" or (tipo == "secao" and b["texto"].strip().upper() in ("RELATÓRIO", "VOTO"))
        p, runs = novo_par(doc, tipo, b["texto"], quebra_antes=quebra)
        if b.get("comentario") and runs:
            pendentes.append((runs, b["comentario"]))
        sep = ESTILOS[tipo]["vazio"]
        if tipo in TRANSCRICOES and prox not in TRANSCRICOES:
            sep = True
        if tipo.startswith(("ementa", "resumo")) and prox == "secao":
            sep = False
        if tipo == "subtitulo" and prox == "subtitulo":
            sep = False
        if tipo == "secao" and b["texto"].strip().upper() == "EMENTA":
            sep = False
        if sep and prox is not None:
            vazio(doc)
    for runs, texto in pendentes:
        doc.add_comment(runs, text=PREFIXO_COMENTARIO + texto, author=AUTOR_COMENTARIO, initials="AG")
    doc.save(saida)
    return len(doc.paragraphs), len(pendentes)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    modelo = next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--modelo=")), MODELO_PADRAO)
    spec = json.load(open(args[0], encoding="utf-8"))
    n, c = montar(spec, args[1], modelo)
    print(f"OK: {args[1]} · {n} parágrafos · {c} comentário(s)")
