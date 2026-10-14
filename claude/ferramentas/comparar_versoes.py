#!/usr/bin/env python3
"""Compara a versão do agente (A) com a versão final (B): python3 comparar_versoes.py AGENTE.docx FINAL.docx
Usa a versão ACEITA de cada arquivo e mede quais alterações rastreadas do agente sobreviveram na final.
Exige revisao_tracked.py na mesma pasta."""
import sys, re, difflib, zipfile, os, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.oxml.ns import qn
from revisao_tracked import aceitar_tudo, contar_alteracoes

def secoes(doc):
    sec, out = "CABEÇALHO", []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.upper() in ("EMENTA", "RELATÓRIO", "VOTO", "DISPOSITIVO"): sec = t.upper()
        if t: out.append((sec, t, p))
    return out

def fmt(p):
    runs = [r for r in p.runs if r.text.strip()]
    return {"alin": str(p.alignment).split(".")[-1].split(" ")[0] if p.alignment is not None else "herdado",
            "negrito": any(r.bold for r in runs), "tam": sorted({r.font.size.pt for r in runs if r.font.size}),
            "realce": any(r.font.highlight_color for r in runs)}

def comentarios(path):
    with zipfile.ZipFile(path) as z:
        if "word/comments.xml" not in z.namelist(): return []
        x = z.read("word/comments.xml").decode("utf8")
    return [re.sub(r"<[^>]+>", "", c).strip()[:160] for c in re.findall(r"<w:comment [^>]*>(.*?)</w:comment>", x, re.S)]

def palavras(t): return len(t.split())
def periodos(t): return len([x for x in re.split(r"(?<=[.!?])\s+(?=[A-ZÀ-Ú])", t) if x.strip()])

def aceito(path):
    c = contar_alteracoes(path)
    if sum(c.values()) == 0:
        return path, c
    out = os.path.join(tempfile.mkdtemp(), "aceito.docx")
    return aceitar_tudo(path, out), c

pa, ca_ = aceito(sys.argv[1]); pb, cb_ = aceito(sys.argv[2])
A, B = secoes(Document(pa)), secoes(Document(pb))
ta, tb = [t for _, t, _ in A], [t for _, t, _ in B]
sm = difflib.SequenceMatcher(None, ta, tb, autojunk=False)

def wdiff(a, b):
    wa, wb = a.split(), b.split(); s = difflib.SequenceMatcher(None, wa, wb); o = []
    for op, i1, i2, j1, j2 in s.get_opcodes():
        if op == "equal": o.append(" ".join(wa[i1:i2]) if i2 - i1 <= 6 else " ".join(wa[i1:i1+3]) + " … " + " ".join(wa[i2-3:i2]))
        if op in ("delete", "replace"): o.append("[-" + " ".join(wa[i1:i2]) + "-]")
        if op in ("insert", "replace"): o.append("{+" + " ".join(wb[j1:j2]) + "+}")
    return " ".join(o)

print("# COMPARAÇÃO AGENTE × FINAL\n")
n = 0
for op, i1, i2, j1, j2 in sm.get_opcodes():
    if op == "equal":
        for k in range(i2 - i1):
            fa, fb = fmt(A[i1+k][2]), fmt(B[j1+k][2])
            dif = {c: (fa[c], fb[c]) for c in fa if fa[c] != fb[c]}
            if dif:
                n += 1; print(f"## [{n}] FORMATAÇÃO · {B[j1+k][0]}\n«{tb[j1+k][:120]}»\n{dif}\n")
        continue
    n += 1; sec = (B[j1][0] if j1 < len(B) else A[i1][0])
    bloco_a, bloco_b = ta[i1:i2], tb[j1:j2]
    tipo = {"replace": "ALTERADO", "delete": "SUPRIMIDO", "insert": "ACRESCENTADO"}[op]
    # detecta cisão/fusão de parágrafos
    if op == "replace" and " ".join(bloco_a).split() == " ".join(bloco_b).split():
        tipo = "CISÃO" if len(bloco_b) > len(bloco_a) else "FUSÃO"
    print(f"## [{n}] {tipo} · {sec} · {len(bloco_a)}→{len(bloco_b)} parágrafo(s)")
    if op == "replace" and len(bloco_a) == len(bloco_b):
        for a, b in zip(bloco_a, bloco_b): print("~ " + wdiff(a, b))
    else:
        for a in bloco_a: print(f"- ({palavras(a)}p/{periodos(a)}per) {a[:400]}")
        for b in bloco_b: print(f"+ ({palavras(b)}p/{periodos(b)}per) {b[:400]}")
    print()

def stats(L):
    v = [t for s, t, _ in L if s == "VOTO" and palavras(t) >= 4 and not t.isupper()]
    if not v: return "sem VOTO"
    w = sorted(palavras(t) for t in v)
    return (f"{len(v)} parágrafos no voto · mediana {w[len(w)//2]} palavras · "
            f"{sum(periodos(t) > 1 for t in v)} com 2+ períodos · {sum(palavras(t) < 15 for t in v)} com <15 palavras")
print("# MÉTRICAS DO VOTO\nAGENTE: " + stats(A) + "\nFINAL:  " + stats(B))
ca, cb = comentarios(sys.argv[1]), comentarios(sys.argv[2])
print(f"\n# COMENTÁRIOS\nagente {len(ca)} · final {len(cb)}")
for c in ca:
    if c not in cb: print("- removido: " + c)
for c in cb:
    if c not in ca: print("+ novo: " + c)
# destino das alterações rastreadas do agente (julgado pelo contexto na frase)
norm = lambda x: re.sub(r"\s+", " ", x).strip()
texto_final = norm(" ".join(tb))
W = lambda t: qn("w:" + t)
def vtexto(el, visao):
    if el.tag == W("r"):
        return "".join(x.text or "" for x in el.iter(W("t")))
    if el.tag == W("ins"):
        return "".join(x.text or "" for x in el.iter(W("t"))) if visao == "aceita" else ""
    if el.tag == W("del"):
        return "".join(x.text or "" for x in el.iter(W("delText"))) if visao == "original" else ""
    return ""
ins_ok, ins_rej, del_ok, del_rej = [], [], [], []
for p in Document(sys.argv[1]).element.body.iter(W("p")):
    filhos = [c for c in p if c.tag in (W("r"), W("ins"), W("del"))]
    for i, el in enumerate(filhos):
        if el.tag not in (W("ins"), W("del")):
            continue
        visao = "aceita" if el.tag == W("ins") else "original"
        alvo = vtexto(el, visao)
        if not re.search(r"\w", alvo):
            continue
        antes = "".join(vtexto(x, visao) for x in filhos[:i])[-15:]
        depois = "".join(vtexto(x, visao) for x in filhos[i + 1:])[:15]
        trecho = norm(antes + alvo + depois)
        presente = trecho in texto_final
        if el.tag == W("ins"):
            (ins_ok if presente else ins_rej).append(norm(alvo))
        else:
            (del_rej if presente else del_ok).append(norm(alvo))
print(f"\n# ALTERAÇÕES RASTREADAS DO AGENTE ({ca_})")
if cb_ and sum(cb_.values()):
    print(f"atenção: a versão final ainda tem alterações pendentes {cb_}; a comparação usa a versão aceita")
print(f"inserções mantidas {len(ins_ok)} · rejeitadas/reescritas {len(ins_rej)} | exclusões mantidas {len(del_ok)} · desfeitas {len(del_rej)}")
for t in ins_rej[:15]: print("- inserção rejeitada: «" + t[:140] + "»")
for t in del_rej[:15]: print("- exclusão desfeita (texto voltou): «" + t[:140] + "»")
print(f"\nTotal de diferenças: {n}")
