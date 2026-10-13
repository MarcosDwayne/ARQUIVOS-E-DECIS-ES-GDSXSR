#!/usr/bin/env python3
import json, csv, os, datetime
BASE = os.environ.get('GDSXSR_BASE', '/home/claude/gdsxsr')
HOJE = datetime.date.today().strftime('%d/%m/%Y')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

L = json.load(open(f'{BASE}/indice.json'))
COLS = ['Nível', 'Classe', 'Processo', 'Resultado', 'Assessor', 'Marcos', 'Marcadores', 'Tipo', 'Matéria',
        'Ramo (ementa)', 'Temas', 'Assunto (título)', 'Cabeçalho da ementa', 'Pasta', 'Duplicado', 'Enviado em', 'Arquivo', 'Link', 'ID']
LARG = [7, 8, 13, 10, 18, 8, 16, 14, 30, 30, 10, 55, 70, 13, 10, 12, 60, 14, 36]
F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF', size=10)
HF = PatternFill('solid', fgColor='1F3864')
N = Font(name=F, size=10)
NIV = {1: 'FCE4D6', 2: 'FFF2CC', 3: 'E2EFDA', 4: 'FFFFFF'}
fino = Side(style='thin', color='D9D9D9')

wb = Workbook()
# ---------- Como usar
ws = wb.active; ws.title = 'Como usar'
txt = [
    ('ÍNDICE DE MODELOS — REVISOR DE VOTOS GDSXSR', Font(name=F, bold=True, size=14, color='1F3864')),
    (f'Gerado em {HOJE} a partir das pastas 01 — MODELOS OURO e 02 — MODELOS GERAIS (APC, AGI, EMD) do Drive: {len(L)} arquivos.', N),
    ('', N),
    ('Para que serve', Font(name=F, bold=True, size=11)),
    ('Escolher, em cada revisão, os 3 a 8 modelos comparáveis sem varrer o Drive. Filtre a aba Índice por Classe → Matéria → Resultado e ordene por Nível.', N),
    ('', N),
    ('Nível (ordem de prioridade das instruções do projeto)', Font(name=F, bold=True, size=11)),
    ('1 — MODELO OURO (pasta 01 — MODELOS OURO)', N),
    ('2 — versão corrigida, alterada, ajustada ou aprovada (marcador no nome do arquivo)', N),
    ('3 — versão com participação de Marcos', N),
    ('4 — demais votos do corpus', N),
    ('', N),
    ('Colunas', Font(name=F, bold=True, size=11)),
    ('Classe, Processo, Resultado (NP, DP, DPP, NC), Assessor, Marcadores e Assunto vêm do nome do arquivo.', N),
    ('Matéria é uma classificação automática por palavras-chave do título e do início da ementa; um voto pode ter mais de uma.', N),
    ('Ramo (ementa) e Cabeçalho da ementa vêm do texto do próprio voto; ficam em branco quando o arquivo não começa pela ementa.', N),
    ('Temas lista os números de Tema citados no título ou no início da ementa. Conferir a vigência na fonte oficial antes de usar.', N),
    ('Duplicado = o mesmo processo aparece em mais de uma pasta (ex.: OURO e GERAIS).', N),
    ('Tipo distingue voto, voto-vista/divergente, decisão e despacho.', N),
    ('', N),
    ('Limitações', Font(name=F, bold=True, size=11)),
    ('A classificação automática pode errar em títulos abreviados. Na dúvida, abra o arquivo pelo link.', N),
    (f'O índice reflete o Drive em {HOJE}. Arquivos enviados depois entram na próxima geração mensal.', N),
    ('O modelo orienta estilo e estrutura; não prova fatos do processo revisado nem torna vigente o direito que cita.', N),
]
for i, (t, f) in enumerate(txt, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 120

# ---------- Índice
wi = wb.create_sheet('Índice')
for j, (c, w) in enumerate(zip(COLS, LARG), 1):
    cell = wi.cell(row=1, column=j, value=c); cell.font = H; cell.fill = HF
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    wi.column_dimensions[get_column_letter(j)].width = w
for i, l in enumerate(L, 2):
    for j, c in enumerate(COLS, 1):
        v = l.get(c, '')
        if c == 'Link':
            cell = wi.cell(row=i, column=j, value='abrir'); cell.hyperlink = v
            cell.font = Font(name=F, size=10, color='0563C1', underline='single')
        else:
            cell = wi.cell(row=i, column=j, value=v); cell.font = N
        cell.alignment = Alignment(vertical='top', wrap_text=c in ('Assunto (título)', 'Matéria'))
        cell.border = Border(bottom=fino)
    wi.cell(row=i, column=1).fill = PatternFill('solid', fgColor=NIV[l['Nível']])
    wi.cell(row=i, column=1).alignment = Alignment(horizontal='center', vertical='top')
fim = len(L) + 1
tab = Table(displayName='Indice', ref=f'A1:{get_column_letter(len(COLS))}{fim}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
wi.add_table(tab)
wi.freeze_panes = 'D2'
wi.row_dimensions[1].height = 30

# ---------- Resumo (fórmulas sobre a aba Índice)
wr = wb.create_sheet('Resumo')
B = Font(name=F, bold=True, size=10)
rng = lambda col: f"'Índice'!${col}$2:${col}${fim}"
colA, colB, colD, colI = 'A', 'B', 'D', 'I'
wr['A1'] = 'RESUMO DO ÍNDICE'; wr['A1'].font = Font(name=F, bold=True, size=13, color='1F3864')
wr['A2'] = 'Contagens calculadas por fórmula a partir da aba Índice.'; wr['A2'].font = N
r = 4
classes = ['APC', 'AGI', 'EMD']
def cab(row, itens):
    for j, t in enumerate(itens, 1):
        c = wr.cell(row=row, column=j, value=t); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center')
wr.cell(row=r, column=1, value='Classe × Nível').font = B; r += 1
cab(r, ['Classe', 'Nível 1', 'Nível 2', 'Nível 3', 'Nível 4', 'Total']); r += 1
ini = r
for cl in classes:
    wr.cell(row=r, column=1, value=cl).font = N
    for k in range(1, 5):
        wr.cell(row=r, column=k + 1, value=f'=COUNTIFS({rng(colB)},$A{r},{rng(colA)},{k})').font = N
    wr.cell(row=r, column=6, value=f'=SUM(B{r}:E{r})').font = B; r += 1
wr.cell(row=r, column=1, value='Outras classes').font = N
for k in range(1, 5):
    wr.cell(row=r, column=k + 1, value=f'=COUNTIFS({rng(colA)},{k})-SUM({get_column_letter(k+1)}{ini}:{get_column_letter(k+1)}{r-1})').font = N
wr.cell(row=r, column=6, value=f'=SUM(B{r}:E{r})').font = B; r += 1
wr.cell(row=r, column=1, value='Total').font = B
for k in range(2, 7):
    col = get_column_letter(k); wr.cell(row=r, column=k, value=f'=SUM({col}{ini}:{col}{r-1})').font = B
r += 2
wr.cell(row=r, column=1, value='Classe × Resultado').font = B; r += 1
cab(r, ['Classe', 'NP', 'DP', 'DPP', 'NC', 'Sem sigla']); r += 1
for cl in classes:
    wr.cell(row=r, column=1, value=cl).font = N
    for k, res in enumerate(['NP', 'DP', 'DPP', 'NC'], 2):
        wr.cell(row=r, column=k, value=f'=COUNTIFS({rng(colB)},$A{r},{rng(colD)},"{res}")').font = N
    wr.cell(row=r, column=6, value=f'=COUNTIFS({rng(colB)},$A{r},{rng(colD)},"")').font = N
    r += 1
r += 1
wr.cell(row=r, column=1, value='Matéria × Classe (um voto pode ter mais de uma matéria)').font = B; r += 1
cab(r, ['Matéria', 'APC', 'AGI', 'EMD', 'Nível 1 (todas as classes)', 'Total']); r += 1
mats = ['Processual', 'Execução e cumprimento', 'Bancário e consumidor', 'Família', 'Responsabilidade civil',
        'Fazenda Pública e servidor', 'Saúde', 'Contratos e obrigações', 'Posse e imóveis', 'Busca e apreensão', 'PASEP']
for m in mats:
    wr.cell(row=r, column=1, value=m).font = N
    for k, cl in enumerate(classes, 2):
        wr.cell(row=r, column=k, value=f'=COUNTIFS({rng(colB)},"{cl}",{rng(colI)},"*"&$A{r}&"*")').font = N
    wr.cell(row=r, column=5, value=f'=COUNTIFS({rng(colA)},1,{rng(colI)},"*"&$A{r}&"*")').font = N
    wr.cell(row=r, column=6, value=f'=COUNTIFS({rng(colI)},"*"&$A{r}&"*")').font = B
    r += 1
wr.column_dimensions['A'].width = 34
for k in 'BCDEF': wr.column_dimensions[k].width = 14
wr.column_dimensions['E'].width = 24

wb.move_sheet('Índice', offset=-1)
wb.active = 0
os.makedirs('/mnt/user-data/outputs', exist_ok=True)
wb.save('/mnt/user-data/outputs/INDICE-MODELOS-GDSXSR.xlsx')
cols = ['Nível','Classe','Processo','Resultado','Assessor','Marcadores','Tipo','Matéria','Ramo (ementa)','Temas','Assunto (título)','Pasta','ID']
with open(f'{BASE}/INDICE-MODELOS-GDSXSR.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(cols); [w.writerow([l[c] for c in cols]) for l in L]
print('ok', len(L), 'linhas')
